import io
import os
import re

import streamlit as st
import google.generativeai as genai
from pypdf import PdfReader
from docx import Document
from fpdf import FPDF

MODEL = "gemini-2.5-flash"

MASTER_PROMPT = """You are a senior Business Analyst who writes Business Requirement Documents (BRDs).

TASK: Write a complete BRD for the project idea below, following the TEMPLATE exactly.

RULES:
1. Reproduce the template's section order, numbering and headings exactly. Do not add, remove or rename sections.
2. Fill every section with specific, realistic content. If information is missing, make a sensible assumption and mark it "(Assumption)".
3. Use Markdown: '#' for the title, '##' for sections, '-' for bullets, '|' tables where the template uses tables.
4. Requirements must be numbered (e.g. FR-01, NFR-01) and testable.
5. Output only the BRD. No preamble, no commentary.
{banking}
=== TEMPLATE ===
{template}

=== PROJECT IDEA ===
{idea}
"""

BANKING_CONTEXT = """
BANKING CONTEXT (apply where relevant):
- Include KYC/AML/CDD requirements and customer onboarding controls.
- Reference State Bank of Pakistan (SBP) regulations, circulars and reporting requirements; mark exact clause numbers as "(to be verified by Compliance)".
- Include audit trail, data privacy, segregation of duties and maker-checker controls.
- Include risk, fraud and business continuity considerations.
"""

DEFAULT_TEMPLATE = """# Business Requirement Document
## 1. Document Control
## 2. Executive Summary
## 3. Business Objectives
## 4. Scope (In Scope / Out of Scope)
## 5. Stakeholders
## 6. Functional Requirements
## 7. Non-Functional Requirements
## 8. Assumptions & Constraints
## 9. Risks & Mitigation
## 10. Success Criteria / KPIs
## 11. Approval
"""


# ---------- helpers ----------
def extract_text(uploaded):
    name = uploaded.name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(uploaded)
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        doc = Document(uploaded)
        parts = []
        for p in doc.paragraphs:
            if not p.text.strip():
                continue
            style = p.style.name.lower()
            if style.startswith("heading"):
                level = re.findall(r"\d+", style)
                parts.append("#" * (int(level[0]) + 1 if level else 2) + " " + p.text)
            else:
                parts.append(p.text)
        for t in doc.tables:
            for row in t.rows:
                parts.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        return "\n".join(parts)
    return ""


def generate_brd(template, idea, banking=False, api_key=None):
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(MODEL)
    prompt = MASTER_PROMPT.format(
        banking=BANKING_CONTEXT if banking else "", template=template, idea=idea
    )
    return model.generate_content(prompt).text


def clean(s):
    return re.sub(r"\*\*|__|`", "", s)


def to_docx(md):
    doc = Document()
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=cols)
                t.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        t.cell(r, c).text = clean(val)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(clean(m.group(2)), level=min(len(m.group(1)) - 1, 3))
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(clean(re.sub(r"^\s*[-*]\s+", "", line)), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(clean(line))
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(md):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    for line in md.splitlines():
        text = clean(line).encode("latin-1", "replace").decode("latin-1")
        m = re.match(r"^(#{1,4})\s+(.*)", text)
        pdf.set_x(pdf.l_margin)
        if m:
            size = {1: 18, 2: 14, 3: 12, 4: 11}[len(m.group(1))]
            pdf.set_font("Helvetica", "B", size)
            pdf.multi_cell(0, 8, m.group(2))
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, text.replace("\t", "    "))
    return bytes(pdf.output())


# ---------- UI ----------
st.set_page_config(page_title="BRD Generator", page_icon="📄")
st.title("BRD Generator")

api_key = st.secrets.get("GEMINI_API_KEY", None) if hasattr(st, "secrets") else None
api_key = api_key or os.getenv("GEMINI_API_KEY")
if not api_key:
    api_key = st.text_input("Gemini API key", type="password")

uploaded = st.file_uploader("1. Upload BRD format (PDF / DOCX)", type=["pdf", "docx"])
template = DEFAULT_TEMPLATE
if uploaded:
    extracted = extract_text(uploaded)
    if extracted.strip():
        template = extracted
        with st.expander("Extracted format (editable)"):
            template = st.text_area("Template", extracted, height=250)
    else:
        st.warning("Couldn't read text from this file (scanned PDF?). Using default format.")
else:
    st.caption("No file uploaded: a standard BRD format will be used.")

idea = st.text_area("2. Project idea", placeholder="Loan app for farmers", height=150)
banking = st.toggle("Banking context (KYC, SBP rules)")

if st.button("Generate BRD", type="primary"):
    if not api_key:
        st.error("Add your Gemini API key.")
    elif not idea.strip():
        st.error("Enter a project idea.")
    else:
        with st.spinner("Writing your BRD..."):
            try:
                st.session_state["brd"] = generate_brd(template, idea, banking, api_key)
            except Exception as e:
                st.error(f"Generation failed: {e}")

if "brd" in st.session_state:
    brd = st.session_state["brd"]
    st.divider()
    st.markdown(brd)
    c1, c2 = st.columns(2)
    c1.download_button(
        "Download DOCX",
        to_docx(brd),
        "BRD.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    c2.download_button("Download PDF", to_pdf(brd), "BRD.pdf", "application/pdf")
